"""
Project Atlas — Hackathon Finals Pitch Deck Generator
scripts/generate_pitch_deck.py

Generates a 16:9 widescreen presentation deck (Atlas_Finals_Pitch_Deck.pptx)
following the exact 10-slide architectural blueprint in docs/final_plan.md.
Adheres to world-class deck standards: high-contrast executive palettes,
bold typography, clear metric callouts, and structured comparison matrices.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Executive Enterprise Color Palette
COLOR_BG_DARK = RGBColor(15, 23, 42)        # Deep Slate #0f172a
COLOR_CARD_DARK = RGBColor(30, 41, 59)      # Slate Card #1e293b
COLOR_ACCENT_BLUE = RGBColor(14, 165, 233)   # Sky Blue #0ea5e9
COLOR_ACCENT_GREEN = RGBColor(16, 185, 129)  # Emerald #10b981
COLOR_TEXT_WHITE = RGBColor(248, 250, 252)   # Near White #f8fafc
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400 #94a3b8
COLOR_BORDER = RGBColor(51, 65, 85)         # Slate 700 #334155
COLOR_ACCENT_RED = RGBColor(239, 68, 68)     # Red 500 #ef4444


def create_deck(output_path: str = "docs/Atlas_Finals_Pitch_Deck.pptx"):
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Blank layout
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color=COLOR_BG_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_DARK, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    def add_header(slide, category: str, title: str, subtitle: str = ""):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.3))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT_BLUE

        # Title
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_WHITE
        p_title.space_before = Pt(4)

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = COLOR_TEXT_MUTED
            p_sub.space_before = Pt(4)

    # =========================================================================
    # SLIDE 1: Title & Positioning (The Vision Exhale)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Center hero card
    add_card(slide1, Inches(1.5), Inches(1.6), Inches(10.33), Inches(4.3), COLOR_CARD_DARK, COLOR_ACCENT_BLUE)

    tx1 = slide1.shapes.add_textbox(Inches(1.8), Inches(2.0), Inches(9.7), Inches(3.5))
    tf1 = tx1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "PROJECT ATLAS"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "The Universal Web Autopilot — Dual-Speed Cognitive Agent"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT_BLUE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Sub-100ms Reflexes • Zero-Knowledge Cryptographic Vault • 100% WCAG 2.2 AAA Accessibility"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(16)

    p4 = tf1.add_paragraph()
    p4.text = "Hackathon Finals Presentation  |  October 2026"
    p4.font.size = Pt(12)
    p4.font.color.rgb = COLOR_ACCENT_GREEN
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(24)

    # =========================================================================
    # SLIDE 2: The Core Problem (The T-Chart Fracture)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "The Industry Dilemma", "Why Modern Web Agents Keep Breaking in Production", "Today's browser assistants are brittle, agonizingly slow, and leak enterprise secrets.")

    # Left: Status Quo Failures (Red)
    add_card(slide2, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8), COLOR_CARD_DARK, COLOR_ACCENT_RED)
    tx2_left = slide2.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.4))
    tf2_l = tx2_left.text_frame
    tf2_l.word_wrap = True

    p = tf2_l.paragraphs[0]
    p.text = "CURRENT BROWSER AGENTS"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_RED

    points_left = [
        ("Fragile Selectors", "Break constantly when React/Vue re-renders DOM or shifts CSS classes."),
        ("Multi-Second Cloud Latency", "3,500ms+ per action step causes jarring freeze-ups and thundering-herd timeouts."),
        ("Catastrophic Data Leakage", "Raw passwords, patient IDs, and credit cards sent unencrypted to cloud LLMs."),
        ("Accessibility Failure", "Ignore screen-readers, lack visible focus states, and violate WCAG compliance.")
    ]
    for title, desc in points_left:
        p_item = tf2_l.add_paragraph()
        p_item.text = f"✖ {title}: "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = COLOR_TEXT_WHITE
        p_item.space_before = Pt(12)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Atlas Architectural Superiority (Green)
    add_card(slide2, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8), COLOR_CARD_DARK, COLOR_ACCENT_GREEN)
    tx2_right = slide2.shapes.add_textbox(Inches(7.2), Inches(2.2), Inches(5.0), Inches(4.4))
    tf2_r = tx2_right.text_frame
    tf2_r.word_wrap = True

    p = tf2_r.paragraphs[0]
    p.text = "THE PROJECT ATLAS STANDARD"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_GREEN

    points_right = [
        ("64-Bit FNV-1a Semantic Fingerprints", "Addresses elements by structural semantics and accessibility trees, not fragile DOM paths."),
        ("Dual-Speed Reflex Engine", "System 1 evaluates micro-actions in <80ms; System 2 plans high-level goals."),
        ("Client-Side Web Crypto Vault", "AES-GCM encrypted zero-knowledge storage; credentials never leave client un-tokenized."),
        ("WCAG 2.2 AAA by Design", "Built-in voice in/out, high-contrast outlines, screen-reader live regions, and notch UI.")
    ]
    for title, desc in points_right:
        p_item = tf2_r.add_paragraph()
        p_item.text = f"✔ {title}: "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = COLOR_TEXT_WHITE
        p_item.space_before = Pt(12)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: TAM & Urgency (The Stat Callout)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Market Urgency & Opportunity", "Massive Commercial TAM Driven by Accessibility & Automation", "The convergence of assistive web compliance and autonomous workforce efficiency.")

    # 3 Stat Cards
    metrics = [
        ("1.3 BILLION", "Global population living with significant disabilities who face exclusionary web barriers daily."),
        ("42% ABANDON", "Average drop-off rate on multi-step enterprise workflows (insurance, healthcare, banking)."),
        ("$13.5 BILLION", "Total addressable annual market for enterprise assistive agents and automated web workflows.")
    ]
    col_w = Inches(3.6)
    for i, (val, caption) in enumerate(metrics):
        col_l = Inches(0.8 + i * 4.0)
        add_card(slide3, col_l, Inches(2.2), col_w, Inches(4.4))
        tx = slide3.shapes.add_textbox(col_l + Inches(0.2), Inches(2.5), col_w - Inches(0.4), Inches(3.8))
        tf = tx.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE
        p.alignment = PP_ALIGN.CENTER

        p_desc = tf.add_paragraph()
        p_desc.text = caption
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = COLOR_TEXT_WHITE
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.space_before = Pt(20)

    # =========================================================================
    # SLIDE 4: Architecture (The Hero Split)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "System Architecture", "Dual-Speed Cognitive Engine & Zero-Trust Client Vault", "Decoupling ultrafast sub-100ms micro-decisions from high-level reasoning and planning.")

    # Left: Mechanical Architecture Steps
    add_card(slide4, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8))
    tx4_l = slide4.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.4))
    tf4_l = tx4_l.text_frame
    tf4_l.word_wrap = True

    p = tf4_l.paragraphs[0]
    p.text = "HOW ATLAS OPERATES"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE

    steps = [
        ("1. Perceive", "DOM Serializer parses accessibility trees and emits stable 64-bit FNV-1a hashes in <16ms."),
        ("2. System 1 Reflex", "TypeSafe Jev 3 evaluates micro-actions in <80ms without cloud roundtrips."),
        ("3. System 2 Reason", "Llama 3.3 / GPT reasoners decompose complex multi-hop goals into milestones."),
        ("4. Zero-Trust Vault", "Client-side Web Crypto AES-GCM tokenizes credentials before transmission.")
    ]
    for title, desc in steps:
        p_s = tf4_l.add_paragraph()
        p_s.text = f"{title}: "
        p_s.font.bold = True
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = COLOR_TEXT_WHITE
        p_s.space_before = Pt(10)
        run = p_s.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Latency & Security Box
    add_card(slide4, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8))
    tx4_r = slide4.shapes.add_textbox(Inches(7.2), Inches(2.2), Inches(5.0), Inches(4.4))
    tf4_r = tx4_r.text_frame
    tf4_r.word_wrap = True

    p = tf4_r.paragraphs[0]
    p.text = "DUAL-SPEED BENCHMARK"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_GREEN

    comparisons = [
        ("Action Latency", "Atlas: ~75ms  |  Standard LLM Agent: ~3,800ms (50x Faster)"),
        ("Token Efficiency", "Compact budget serializer cuts DOM token consumption by 82%"),
        ("Secret Storage", "AES-256-GCM in IndexedDB (Client-only, non-exportable key)"),
        ("Resilience Quorum", "Automated 429 backoff, semaphore throttling & speculative cancellation")
    ]
    for metric, val in comparisons:
        p_c = tf4_r.add_paragraph()
        p_c.text = f"• {metric}:"
        p_c.font.bold = True
        p_c.font.size = Pt(13)
        p_c.font.color.rgb = COLOR_TEXT_WHITE
        p_c.space_before = Pt(10)
        run = p_c.add_run()
        run.text = f"\n  {val}"
        run.font.bold = False
        run.font.color.rgb = COLOR_ACCENT_BLUE

    # =========================================================================
    # SLIDE 5: Demonstration Flow Rail
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "End-to-End Execution", "Autonomous Multi-Hop Flow Rail", "Live Scenario: 'Sign up for university portal, complete medical triage, and download timetable.'")

    # 4 Sequential Milestones
    flow_steps = [
        ("1. Page Classification", "Scout identifies form structure and classifies page semantics in <40ms."),
        ("2. Tokenized Autofill", "Vault injects {profile.email} and {password} without plaintext leaks."),
        ("3. Human-in-Loop Gate", "Accessible confirmation modal halts execution for user sign-off."),
        ("4. File Download & Verify", "Executor clicks timetable download and verifier quorum confirms delivery.")
    ]
    step_w = Inches(2.7)
    for i, (title, desc) in enumerate(flow_steps):
        left_pos = Inches(0.8 + i * 3.0)
        add_card(slide5, left_pos, Inches(2.3), step_w, Inches(4.2))
        tx = slide5.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.5), step_w - Inches(0.3), Inches(3.8))
        tf = tx.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = f"STAGE 0{i+1}"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_BLUE

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE
        p_t.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(12)

    # =========================================================================
    # SLIDE 6: Competitive Landscape
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Competitive Landscape", "The Unfair Advantage: Speed + Zero-Trust Privacy", "Where Project Atlas sits compared to OpenAI Operator, Anthropic Computer Use, and Puppeteer.")

    add_card(slide6, Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
    tx6 = slide6.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(11.0), Inches(4.2))
    tf6 = tx6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "COMPETITIVE BENCHMARK MATRIX"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE

    matrix = [
        ("OpenAI Operator / Anthropic", "High reasoning capability, but 4,000ms latency, high cost, and zero client privacy vault."),
        ("Raw Puppeteer / Selenium", "Fast execution, but brittle CSS selectors, zero AI adaptation, and breaks on layout shifts."),
        ("Traditional Screen Readers", "Provides auditory output, but requires tedious manual keyboard stepping with zero autonomy."),
        ("Project Atlas (The Sweet Spot)", "Sub-100ms reflex speed + System 2 planning + Web Crypto zero-knowledge vault + WCAG 2.2 AAA.")
    ]
    for player, desc in matrix:
        p_m = tf6.add_paragraph()
        p_m.text = f"• {player}: "
        p_m.font.bold = True
        p_m.font.size = Pt(13)
        p_m.font.color.rgb = COLOR_TEXT_WHITE if "Atlas" not in player else COLOR_ACCENT_GREEN
        p_m.space_before = Pt(12)
        run = p_m.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Technical Rigor & Quality Gates
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Technical Rigor & Verification", "Tested for Extreme Production Environments", "A fully green test pyramid backed by deterministic replay harnesses.")

    rigor_cards = [
        ("SECURITY & PRIVACY", "• Web Crypto AES-GCM vault\n• Zero plaintext secret leaks\n• Strict Manifest V3 CSP\n• Isolated Closed Shadow DOM"),
        ("AUTOMATED TEST PYRAMID", "• 477+ backend pytest suite\n• Deterministic replay runner\n• Playwright E2E browser harness\n• 100% green pass rate"),
        ("ENTERPRISE RESILIENCE", "• Concurrency semaphores\n• Automatic 429 backoff & retry\n• Speculative task cancellation\n• Arabic/Unicode localization")
    ]
    for i, (title, bullets) in enumerate(rigor_cards):
        left_pos = Inches(0.8 + i * 4.0)
        add_card(slide7, left_pos, Inches(2.2), Inches(3.6), Inches(4.5))
        tx = slide7.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.5), Inches(3.2), Inches(4.0))
        tf = tx.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT_BLUE

        p_b = tf.add_paragraph()
        p_b.text = bullets
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(16)

    # =========================================================================
    # SLIDE 8: Universal Design & Accessibility (WCAG 2.2 AAA)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Accessibility Supremacy", "Built for Universal Independence (WCAG 2.2 AAA)", "Accessibility is the core architectural foundation of Atlas, not an afterthought.")

    add_card(slide8, Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
    tx8 = slide8.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(11.0), Inches(4.2))
    tf8 = tx8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "ACCESSIBILITY FEATURES BUILT INTO EVERY HOP"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE

    a11y_features = [
        ("Multimodal Speech Engine", "Bidirectional voice in/out powered by Web Speech API and streaming TTS synthesis."),
        ("AAA Color Contrast & Focus Outlines", "High-contrast visual targets (>7:1 ratio) with persistent visible focus rings."),
        ("Draggable Floating Notch UI", "Non-intrusive floating pill with pointer capture and Euclidean click discrimination."),
        ("Screen-Reader ARIA Live Regions", "Real-time step updates announced without disrupting native host page navigation.")
    ]
    for feat, desc in a11y_features:
        p_f = tf8.add_paragraph()
        p_f.text = f"✔ {feat}: "
        p_f.font.bold = True
        p_f.font.size = Pt(13)
        p_f.font.color.rgb = COLOR_TEXT_WHITE
        p_f.space_before = Pt(12)
        run = p_f.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 9: Business Model & Go-To-Market
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Commercialization & Growth", "Dual-Engine Revenue Strategy (B2C + Enterprise B2B)", "Capturing both consumer assistive demand and enterprise back-office automation budgets.")

    # Left: B2C Freemium
    add_card(slide9, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8))
    tx9_l = slide9.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.4))
    tf9_l = tx9_l.text_frame
    tf9_l.word_wrap = True

    p = tf9_l.paragraphs[0]
    p.text = "B2C ASSISTIVE FREEMIUM"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE

    p1 = tf9_l.add_paragraph()
    p1.text = "• Free Tier: Unlimited voice & keyboard navigation on all standard websites."
    p1.font.size = Pt(13)
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_before = Pt(12)

    p2 = tf9_l.add_paragraph()
    p2.text = "• Pro Tier ($9/mo): Multi-hop autonomous workflows, priority cloud reasoning, and encrypted profile syncing."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.space_before = Pt(10)

    p3 = tf9_l.add_paragraph()
    p3.text = "• Channel: Direct Chrome Web Store installation with 1-click onboarding."
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.space_before = Pt(10)

    # Right: B2B Enterprise
    add_card(slide9, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8))
    tx9_r = slide9.shapes.add_textbox(Inches(7.2), Inches(2.2), Inches(5.0), Inches(4.4))
    tf9_r = tx9_r.text_frame
    tf9_r.word_wrap = True

    p = tf9_r.paragraphs[0]
    p.text = "B2B ENTERPRISE WORKFORCE"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_GREEN

    p1 = tf9_r.add_paragraph()
    p1.text = "• Per-Seat License ($39/seat/mo): Automated data entry into legacy ERPs, CRMs, and EHR health portals."
    p1.font.size = Pt(13)
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_before = Pt(12)

    p2 = tf9_r.add_paragraph()
    p2.text = "• Zero-Knowledge Guarantee: Critical for healthcare (HIPAA), banking (SOC-2), and government compliance."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.space_before = Pt(10)

    p3 = tf9_r.add_paragraph()
    p3.text = "• Deployment: Dockerized on-premise container stack with tenant audit logging."
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.space_before = Pt(10)

    # =========================================================================
    # SLIDE 10: The Vision Exhale & Call to Action
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)

    add_card(slide10, Inches(1.5), Inches(1.6), Inches(10.33), Inches(4.3), COLOR_CARD_DARK, COLOR_ACCENT_GREEN)
    tx10 = slide10.shapes.add_textbox(Inches(1.8), Inches(2.0), Inches(9.7), Inches(3.5))
    tf10 = tx10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = '"The future of the web isn\'t typing and clicking.\nIt\'s asking and arriving."'
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

    p2 = tf10.add_paragraph()
    p2.text = "WELCOME TO PROJECT ATLAS"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT_GREEN
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)

    p3 = tf10.add_paragraph()
    p3.text = "Live Demo: http://localhost:5500  |  GitHub: Abdelrahman-Aboelnour06/Project-Atlas"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(16)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"[OK] 16:9 Presentation Deck generated successfully: {output_path}")


if __name__ == "__main__":
    create_deck()
